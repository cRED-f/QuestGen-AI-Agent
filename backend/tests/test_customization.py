"""End-to-end customization route tests, no live LLM provider required."""
import io
import json
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from backend import app as module, customization, generation


@pytest.fixture
def ctx(tmp_path, monkeypatch):
    monkeypatch.setenv('QUESTGEN_DB_PATH', str(tmp_path / 'test.sqlite3'))
    with TestClient(module.app) as client:
        course = client.get('/api/courses').json()[0]['id']
        thread = client.post('/api/threads', json={'course_id': course}).json()['id']
        job = module.identifier()
        blueprint = generation.AssessmentBlueprint(title='Data Structures Exam', total_marks=10,
            questions=[generation.QuestionPlan(topic='binary search', kind='mcq', difficulty='medium',
                bloom_level='apply', marks=5), generation.QuestionPlan(topic='linked lists',
                kind='short_answer', difficulty='easy', bloom_level='understand', marks=5)])
        with module.database() as conn:
            t = module.now()
            conn.execute('''INSERT INTO assessment_jobs(id,thread_id,course_id,prompt,provider,model,status,stage,blueprint,created_at,updated_at)
                VALUES (?,?,?,?,?,?,'completed','completed',?,?,?)''',
                (job, thread, course, 'Create two questions', 'ollama', 'test-model', blueprint.model_dump_json(), t, t))
            for i, plan in enumerate(blueprint.questions, start=1):
                content = dict(plan.model_dump(),
                    stem=f'Explain how {plan.topic} works and provide a useful real-world example.',
                    options=['A sorted array', 'An unsorted heap', 'An unordered set', 'A random graph'] if i == 1 else [],
                    correct_answer='A sorted array' if i == 1 else 'A linked list uses linked nodes.',
                    explanation='The correct answer follows from the underlying data structure.',
                    source_ids=[], sources=[], grounded=False)
                conn.execute('''INSERT INTO assessment_questions(id,job_id,ordinal,content,created_at)
                    VALUES(?,?,?,?,?)''', (f'q{i}', job, i, json.dumps(content), t))
                conn.execute('''INSERT INTO assessment_question_versions(question_id,version,content,change_type,instruction,created_at)
                    VALUES(?,1,?,'generated','Initial question',?)''', (f'q{i}', json.dumps(content), t))
        prefix = f'/api/threads/{thread}/generations/{job}'
        yield client, prefix, job


def test_edit_versions_locks_and_conflicts(ctx):
    client, prefix, _ = ctx
    qpath = f'{prefix}/questions/q1'
    initial = client.get(qpath + '/versions').json()
    original = initial['content']
    revised = {k: original[k] for k in ('stem', 'options', 'correct_answer', 'explanation')}
    revised['stem'] = 'Explain binary search using a sorted collection with five elements.'
    result = client.put(qpath, json={'expected_version': 1, **revised})
    assert result.status_code == 200, result.text
    assert result.json()['version'] == 2
    assert len(client.get(qpath + '/versions').json()['versions']) == 2
    assert client.put(qpath, json={'expected_version': 1, **revised}).status_code == 409
    assert client.put(qpath + '/lock?locked=true', json={'expected_version': 2}).json()['locked']
    assert client.put(qpath, json={'expected_version': 2, **revised}).status_code == 409
    assert client.put(qpath + '/lock?locked=false', json={'expected_version': 2}).json()['locked'] is False
    invalid = dict(revised, options=['same', 'same', 'third', 'fourth'])
    assert client.put(qpath, json={'expected_version': 2, **invalid}).status_code == 422
    assert client.get(qpath + '/versions').json()['version'] == 2


def test_review_and_student_export_gate(ctx):
    client, prefix, _ = ctx
    student = prefix + '/export/student.pdf'
    assert client.get(student).status_code == 409
    assert client.put(prefix + '/questions/q1/review', json={'expected_version': 1, 'decision':'approved'}).status_code == 200
    assert client.get(student).status_code == 409
    assert client.put(prefix + '/questions/q2/review', json={'expected_version': 1, 'decision':'approved'}).status_code == 200
    pdf = client.get(student)
    assert pdf.status_code == 200
    assert pdf.headers['content-type'].startswith('application/pdf')
    reader = PdfReader(io.BytesIO(pdf.content))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert 'binary search' in text
    assert 'Answer:' not in text and 'Explanation:' not in text
    assert 'A linked list uses linked nodes.' not in text
    instructor_pdf = client.get(prefix + '/export/instructor.pdf')
    assert 'Answer:' in PdfReader(io.BytesIO(instructor_pdf.content)).pages[0].extract_text()
    student_docx = client.get(prefix + '/export/student.docx')
    document = Document(io.BytesIO(student_docx.content))
    combined = '\n'.join(p.text for p in document.paragraphs)
    assert 'Answer:' not in combined and 'A linked list uses linked nodes.' not in combined
    instructor = Document(io.BytesIO(client.get(prefix + '/export/instructor.docx').content))
    assert 'Answer:' in '\n'.join(p.text for p in instructor.paragraphs)
    # Any edit invalidates approval and blocks student export again.
    original = client.get(prefix + '/questions/q2/versions').json()['content']
    change = {k: original[k] for k in ('stem', 'options', 'correct_answer', 'explanation')}
    change['stem'] = 'Explain how a linked list stores nodes and how traversal operates.'
    assert client.put(prefix + '/questions/q2', json={'expected_version': 1, **change}).status_code == 200
    assert client.get(student).status_code == 409
    assert client.get(prefix + '/questions/q2/versions').json()['review_status'] == 'pending'


def test_reject_unlock_and_audit(ctx):
    client, prefix, _ = ctx
    q = prefix + '/questions/q1'
    result = client.put(q + '/review', json={
        'expected_version':1, 'decision':'rejected','note':'Not suitable for midterm'})
    assert result.status_code == 200
    assert result.json()['review_note'] == 'Not suitable for midterm'
    with module.database() as conn:
        actions = [r['action'] for r in conn.execute('SELECT action FROM assessment_question_audit WHERE question_id=?', ('q1',))]
        assert 'rejected' in actions


def test_targeted_ai_regeneration_is_bounded_and_atomic(ctx, monkeypatch):
    client, prefix, _ = ctx
    count = {'calls':0}
    async def model_json(*args, **kwargs):
        count['calls'] += 1
        return { 'stem': 'What is the best prerequisite for binary search in a collection?',
                'options': ['Sorted array', 'Linked nodes', 'Random graph', 'Heap only'],
                'correct_answer':'Sorted array', 'explanation':'Sorted input enables correct binary search.',
                'source_ids': [] }
    async def review(*args, **kwargs):
        return generation.QualityReview(passed=True)
    monkeypatch.setattr(generation, 'model_json', model_json)
    monkeypatch.setattr(generation, 'review_question', review)
    endpoint = prefix + '/questions/q1/regenerate'
    result = client.post(endpoint, json={'expected_version':1,'instruction':'Use an applied reasoning scenario'})
    assert result.status_code == 200, result.text
    assert result.json()['version'] == 2
    assert count['calls'] == 1
    assert client.get(prefix + '/questions/q2/versions').json()['version'] == 1
    assert client.post(endpoint, json={'expected_version':1,'instruction':'try again'}).status_code == 409
    assert client.put(prefix + '/questions/q1/lock', json={'expected_version':2}).status_code == 200
    assert client.post(endpoint, json={'expected_version':2,'instruction':'try again'}).status_code == 409


def test_conversation_targets_specific_question(ctx, monkeypatch):
    client, prefix, _ = ctx
    async def mock_model_json(settings, system, request):
        return {'actions':[{'ordinal':2,'operation':'lock','instruction':''}]}
    monkeypatch.setattr(generation, 'model_json', mock_model_json)
    result = client.post(prefix + '/customize', json={'instruction':'Lock question 2 only, leave the rest alone'})
    assert result.status_code == 200, result.text
    assert result.json()['actions'] == [{'ordinal':2,'operation':'lock','version':1,'review_status':'pending'}]
    assert client.get(prefix + '/questions/q1/versions').json()['locked'] is False
    assert client.get(prefix + '/questions/q2/versions').json()['locked'] is True
    thread_messages = client.get(prefix.split('/generations/')[0] + '/messages').json()
    assert any('Lock question 2' in m['content'] for m in thread_messages)


def test_course_and_thread_scoped_routes(ctx):
    client, prefix, job = ctx
    wrong = f'/api/threads/nonexistent/generations/{job}'
    assert client.get(wrong + '/export/instructor.pdf').status_code == 404
    assert client.put(wrong + '/questions/q1/review',json={'expected_version':1,'decision':'approved'}).status_code == 404


def test_restore_from_history_creates_new_version(ctx):
    client, prefix, _ = ctx
    endpoint = prefix + '/questions/q1'
    prior = client.get(endpoint + '/versions').json()['content']
    body = {k: prior[k] for k in ('stem', 'options', 'correct_answer', 'explanation')}
    body['stem'] = 'Explain why a sorted array is necessary for binary search algorithms.'
    assert client.put(endpoint, json={**body, 'expected_version':1}).status_code == 200
    restored = client.post(endpoint + '/restore', json={'expected_version':2, 'target_version':1})
    assert restored.status_code == 200
    assert restored.json()['version'] == 3
    assert restored.json()['content']['stem'] == prior['stem']
    assert [v['version'] for v in restored.json()['versions']] == [3,2,1]
    assert client.post(endpoint + '/restore', json={'expected_version':2, 'target_version':1}).status_code == 409


def test_conversation_actions_commit_atomically(ctx, monkeypatch):
    client, prefix, _ = ctx
    async def mock_plan(*args, **kwargs):
        return {'actions': [
            {'ordinal':1,'operation':'approve','instruction':''},
            {'ordinal':2,'operation':'lock','instruction':''},
        ]}
    monkeypatch.setattr(generation, 'model_json', mock_plan)
    result = client.post(prefix + '/customize', json={'instruction':'Approve question 1 and lock question 2'})
    assert result.status_code == 200
    assert client.get(prefix + '/questions/q1/versions').json()['review_status'] == 'approved'
    assert client.get(prefix + '/questions/q2/versions').json()['locked'] is True


def _material_for_revision(client, course_id: str, name: str, body: str) -> str:
    response = client.post(
        f'/api/courses/{course_id}/documents',
        files={'file': (name, body.encode('utf-8'), 'text/plain')},
    )
    assert response.status_code == 201, response.text
    document_id = response.json()['id']
    for _ in range(200):
        item = client.get(f'/api/courses/{course_id}/documents/{document_id}').json()
        if item['status'] in ('ready', 'failed'):
            assert item['status'] == 'ready', item
            return document_id
        time.sleep(0.01)
    pytest.fail('Document indexing timed out')


def test_ai_question_regeneration_respects_attached_assessment_materials(ctx, monkeypatch):
    client, prefix, job_id = ctx
    with module.database() as conn:
        course_id = conn.execute('SELECT course_id FROM assessment_jobs WHERE id=?', (job_id,)).fetchone()['course_id']
    selected = _material_for_revision(client, course_id, 'selected.txt',
        'Binary search of sorted arrays uses a midpoint and repeated halving. ' * 5)
    excluded = _material_for_revision(client, course_id, 'excluded.txt',
        'Binary search in sorted sequences runs in logarithmic time. ' * 5)
    foreign_course = client.post('/api/courses', json={'name': 'Separate course'}).json()['id']
    foreign = _material_for_revision(client, foreign_course, 'foreign.txt',
        'Binary search works on sorted arrays in this unrelated course. ' * 5)
    with module.database() as conn:
        conn.execute('UPDATE assessment_jobs SET document_ids=? WHERE id=?',
                     (json.dumps([selected]), job_id))
    evidence_seen = []

    async def fake_model(settings, system, payload):
        evidence_seen.extend(payload['evidence'])
        return {
            'stem': 'Which requirement allows binary search to locate an element efficiently?',
            'options': ['Sorted array', 'Unordered collection', 'Linked list only', 'Randomized order'],
            'correct_answer': 'Sorted array',
            'explanation': 'A sorted array supports ordered midpoint comparisons and repeated halving.',
            'source_ids': [payload['evidence'][0]['id']],
        }

    async def review(*args, **kwargs):
        return generation.QualityReview(passed=True)

    monkeypatch.setattr(generation, 'model_json', fake_model)
    monkeypatch.setattr(generation, 'review_question', review)
    response = client.post(prefix + '/questions/q1/regenerate', json={
        'expected_version': 1, 'instruction': 'Rewrite with a practical example',
    })
    assert response.status_code == 200, response.text
    saved_sources = response.json()['content']['sources']
    assert {source['document_id'] for source in saved_sources} == {selected}
    assert evidence_seen
    assert excluded not in {source['document_id'] for source in saved_sources}
    assert foreign not in {source['document_id'] for source in saved_sources}


def test_ai_revision_rejects_deleted_selected_document_before_model_call(ctx, monkeypatch):
    client, prefix, job_id = ctx
    with module.database() as conn:
        course_id = conn.execute('SELECT course_id FROM assessment_jobs WHERE id=?', (job_id,)).fetchone()['course_id']
    selected = _material_for_revision(client, course_id, 'missing-after-delete.txt',
        'Binary search uses midpoint comparisons in sorted collections.')
    with module.database() as conn:
        conn.execute('UPDATE assessment_jobs SET document_ids=? WHERE id=?',
                     (json.dumps([selected]), job_id))
    assert client.delete(f'/api/courses/{course_id}/documents/{selected}').status_code == 204

    async def unexpected_model(*args, **kwargs):
        pytest.fail('Model must not run when an assessment source is unavailable')

    monkeypatch.setattr(generation, 'model_json', unexpected_model)
    response = client.post(prefix + '/questions/q1/regenerate', json={
        'expected_version': 1, 'instruction': 'Use a different wording',
    })
    assert response.status_code == 409
    assert client.get(prefix + '/questions/q1/versions').json()['version'] == 1


def test_original_citations_cannot_override_assessment_attachment_filter(ctx, monkeypatch):
    client, prefix, job_id = ctx
    with module.database() as conn:
        course_id = conn.execute('SELECT course_id FROM assessment_jobs WHERE id=?', (job_id,)).fetchone()['course_id']
    selected = _material_for_revision(client, course_id, 'selected-unrelated.txt',
        'Photosynthesis absorbs sunlight through chlorophyll pigment.')
    excluded = _material_for_revision(client, course_id, 'old-citation.txt',
        'Binary search requires a sorted collection and checks the midpoint.')
    with module.database() as conn:
        chunk_id = conn.execute('SELECT id FROM document_chunks WHERE document_id=?', (excluded,)).fetchone()['id']
        old = json.loads(conn.execute('SELECT content FROM assessment_questions WHERE id=?', ('q1',)).fetchone()['content'])
        old['source_ids'] = [chunk_id]
        old['sources'] = [{'chunk_id': chunk_id, 'document_id': excluded, 'snippet': 'Wrong attachment'}]
        conn.execute('UPDATE assessment_questions SET content=? WHERE id=?', (json.dumps(old), 'q1'))
        conn.execute('UPDATE assessment_jobs SET document_ids=? WHERE id=?',
                     (json.dumps([selected]), job_id))
    evidence_seen = []

    async def fake_model(settings, system, payload):
        assert chunk_id not in json.dumps(payload)
        assert excluded not in json.dumps(payload)
        evidence_seen.extend(payload['evidence'])
        return {
            'stem': 'What does a binary search algorithm need before it can compare midpoints?',
            'options': ['Sorted input', 'Random order', 'An unbalanced tree', 'No arrangement'],
            'correct_answer': 'Sorted input',
            'explanation': 'Binary search only works reliably with ordered input for midpoint checks.',
            'source_ids': [],
        }

    async def review(*args, **kwargs):
        return generation.QualityReview(passed=True)

    monkeypatch.setattr(generation, 'model_json', fake_model)
    monkeypatch.setattr(generation, 'review_question', review)
    result = client.post(prefix + '/questions/q1/revise', json={
        'expected_version': 1, 'instruction': 'Clarify the precondition',
    })
    assert result.status_code == 200, result.text
    assert evidence_seen == []
    assert result.json()['content']['source_ids'] == []
    assert result.json()['content']['sources'] == []


def test_conversational_revision_uses_the_same_assessment_source_scope(ctx, monkeypatch):
    client, prefix, job_id = ctx
    with module.database() as conn:
        course_id = conn.execute('SELECT course_id FROM assessment_jobs WHERE id=?', (job_id,)).fetchone()['course_id']
    selected = _material_for_revision(client, course_id, 'lecture-selected.txt',
        'Binary search requires sorted data and repeatedly compares the middle element.')
    excluded = _material_for_revision(client, course_id, 'lecture-excluded.txt',
        'Binary search usually starts by checking the midpoint of sorted input.')
    with module.database() as conn:
        conn.execute('UPDATE assessment_jobs SET document_ids=? WHERE id=?',
                     (json.dumps([selected]), job_id))
    references = []

    async def fake_model(settings, system, request):
        if system.startswith('Translate professor instructions'):
            return {'actions': [{'ordinal': 1, 'operation': 'revise', 'instruction': 'Add practical context'}]}
        references.extend(request['evidence'])
        return {
            'stem': 'Which property of an array lets binary search discard half its elements?',
            'options': ['Sorted elements', 'Unsorted elements', 'Random pointers', 'Mixed data types'],
            'correct_answer': 'Sorted elements',
            'explanation': 'Sorted input makes midpoint comparisons useful for narrowing the search.',
            'source_ids': [request['evidence'][0]['id']],
        }

    async def review(*args, **kwargs):
        return generation.QualityReview(passed=True)

    monkeypatch.setattr(generation, 'model_json', fake_model)
    monkeypatch.setattr(generation, 'review_question', review)
    response = client.post(prefix + '/customize', json={
        'instruction': 'Revise question 1 to include a practical context',
    })
    assert response.status_code == 200, response.text
    assert len(references) == 1
    with module.database() as conn:
        selected_chunk = conn.execute('SELECT id FROM document_chunks WHERE document_id=?', (selected,)).fetchone()['id']
        excluded_chunk = conn.execute('SELECT id FROM document_chunks WHERE document_id=?', (excluded,)).fetchone()['id']
    assert references[0]['id'] == selected_chunk
    assert excluded_chunk not in {r['id'] for r in references}
    revised = client.get(prefix + '/questions/q1/versions').json()['content']
    assert {s['document_id'] for s in revised['sources']} == {selected}
