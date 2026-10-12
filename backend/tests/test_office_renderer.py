"""Prevent false 'install LibreOffice' errors on Windows and conversion failures."""
from __future__ import annotations

import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from backend import ingestion, office_renderer


def test_configured_windows_path_with_spaces_and_quotes(monkeypatch, tmp_path):
    command = tmp_path / "Program Files" / "LibreOffice" / "program" / "soffice.exe"
    command.parent.mkdir(parents=True)
    command.touch()
    monkeypatch.setenv("QUESTGEN_LIBREOFFICE_CMD", f'"{command}"')
    found = office_renderer.locate_office()
    assert found.path == str(command)
    assert found.source == "QUESTGEN_LIBREOFFICE_CMD"


def test_windows_program_files_detected_without_path(monkeypatch, tmp_path):
    command = tmp_path / "Program Files" / "LibreOffice" / "program" / "soffice.exe"
    command.parent.mkdir(parents=True)
    command.touch()
    monkeypatch.delenv("QUESTGEN_LIBREOFFICE_CMD", raising=False)
    monkeypatch.setenv("ProgramFiles", str(tmp_path / "Program Files"))
    monkeypatch.setattr(office_renderer.sys, "platform", "win32")
    monkeypatch.setattr(office_renderer.shutil, "which", lambda _: None)
    result = office_renderer.locate_office()
    assert result.path == str(command)
    assert result.source == "Windows installation"


def test_invalid_explicit_path_reports_configuration_error_not_install_missing(monkeypatch):
    monkeypatch.setenv("QUESTGEN_LIBREOFFICE_CMD", "/not/a/real/soffice")
    found = office_renderer.locate_office()
    assert found.path is None
    assert "QUESTGEN_LIBREOFFICE_CMD" in found.issue


def test_status_checks_executable_launch_not_only_path(monkeypatch):
    monkeypatch.setattr(office_renderer, "locate_office", lambda: office_renderer.OfficeExecutable("soffice", "PATH", None))
    monkeypatch.setattr(office_renderer.subprocess, "run", lambda *a, **kw: SimpleNamespace(returncode=7))
    ready, message, origin = office_renderer.check_office()
    assert ready is False
    assert "found but could not start" in message
    assert origin == "PATH"


def test_conversion_failure_reports_failure_instead_of_install(monkeypatch, tmp_path):
    monkeypatch.setattr(office_renderer, "locate_office", lambda: office_renderer.OfficeExecutable("soffice", "PATH", None))
    monkeypatch.setattr(office_renderer.subprocess, "run", lambda *a, **kw: SimpleNamespace(returncode=42))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Law of supply"
    path = tmp_path / "chapter.pptx"
    deck.save(path)
    with pytest.raises(office_renderer.OfficeRenderError, match="exit code 42"):
        office_renderer.convert_pptx_to_pdf(path, tmp_path)


def test_ingestion_keeps_native_text_and_exposes_render_error(monkeypatch, tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "PPF and opportunity cost"
    path = tmp_path / "chapter.pptx"
    deck.save(path)
    def broken(*args):
        raise office_renderer.OfficeRenderError("LibreOffice failed converting the PowerPoint (exit code 99).")
    monkeypatch.setattr(ingestion, "_render_office", broken)
    dummy_vision = SimpleNamespace(enabled=True, disabled_reason=None, calls=0, max_calls=20, warning=lambda: None)
    warnings = []
    units, pages = ingestion._pptx(path, dummy_vision, warnings)
    assert pages == 1 and len(units) == 1
    assert "PPF and opportunity cost" in units[0].content
    assert "exit code 99" in units[0].issue
    assert "exit code 99" in warnings[0]


def test_real_libreoffice_conversion_when_available(tmp_path):
    ready, _, _ = office_renderer.check_office()
    if not ready:
        pytest.skip("LibreOffice not installed in this test environment")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Chapter 2: Production possibilities"
    path = tmp_path / "chapter.pptx"
    deck.save(path)
    import fitz
    rendered = office_renderer.convert_pptx_to_pdf(path, tmp_path)
    with fitz.open(str(rendered)) as pdf:
        assert len(pdf) == 1
        assert "Production possibilities" in pdf[0].get_text()
