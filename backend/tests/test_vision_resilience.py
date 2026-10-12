"""Provider failures must never erase successful slide analyses or starve later slides."""
from __future__ import annotations

import io
import re
from threading import Event
from pathlib import Path

import fitz
import httpx
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend import app as service, ingestion, vision


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("QUESTGEN_DB_PATH", str(tmp_path / "db.sqlite3"))
    monkeypatch.setenv("QUESTGEN_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("QUESTGEN_VISION_ENABLED", "true")
    monkeypatch.delenv("QUESTGEN_EMBED_MODEL", raising=False)
    monkeypatch.setattr(vision.time, "sleep", lambda _seconds: None)
    with TestClient(service.app) as api:
        response = api.put("/api/settings", json={"provider": "ollama", "model": "test-vision"})
        assert response.status_code == 200
        yield api


def _ppt33() -> bytes:
    pres = Presentation()
    for i in range(1, 34):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(3))
        textbox.text = f"Chapter two slide {i}: opportunity cost and scarcity"
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def _mock_office(source: Path, directory: Path) -> Path:
    result = directory / f"{source.stem}.pdf"
    doc = fitz.open()
    for i in range(33):
        doc.new_page(width=720, height=405).insert_text((50, 50), f"Slide {i+1}")
    doc.save(result)
    doc.close()
    return result


def _wait(client, course, docid):
    for _ in range(200):
        item = client.get(f"/api/courses/{course}/documents/{docid}").json()
        if item["status"] in {"ready", "failed"}:
            return item
        Event().wait(.03)
    raise AssertionError("Ingestion did not finish")


def test_malformed_27th_slide_does_not_skip_slides_28_to_33_and_reprocess_uses_cache(client, monkeypatch):
    monkeypatch.setattr(ingestion, "_render_office", _mock_office)
    fail_slide = {27}
    attempts: list[tuple[int, int]] = []

    class Response:
        def __init__(self, item):
            self.item = item
        def raise_for_status(self):
            pass
        def json(self):
            return self.item

    class Client:
        def __init__(self, *args, **kwargs):
            pass
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def post(self, url, *, json, **kwargs):
            question = json["messages"][1]["content"][0]["text"]
            number = int(re.search(r"slide (\d+)", question).group(1))
            attempts.append((number, json["max_tokens"]))
            text = "" if number in fail_slide else f"Visual transcript for slide {number}"
            return Response({"choices": [{"message": {"content": text}, "finish_reason": "stop"}]})

    monkeypatch.setattr(vision.httpx, "Client", Client)
    cid = client.get("/api/courses").json()[0]["id"]
    result = client.post(f"/api/courses/{cid}/documents", files={"file": ("chapter02.pptx", _ppt33())})
    assert result.status_code == 201
    docid = result.json()["id"]
    document = _wait(client, cid, docid)
    coverage = client.get(f"/api/courses/{cid}/documents/{docid}/coverage").json()
    assert document["status"] == "ready"
    assert coverage["vision"] == 32
    assert [unit["page"] for unit in coverage["units"] if not unit["has_vision"]] == [27]
    assert "Missing slide(s): 27" in document["warning"]
    assert 28 in [slide for slide, _ in attempts] and 33 in [slide for slide, _ in attempts]
    assert [tokens for slide, tokens in attempts if slide == 27] == [2400, 4800, 4800]

    # Recover model service and explicitly reprocess. All 32 successes remain
    # in the per-slide cache and must not incur another provider request.
    before = len(attempts)
    fail_slide.clear()
    assert client.post(f"/api/courses/{cid}/documents/{docid}/reprocess").status_code == 202
    document = _wait(client, cid, docid)
    coverage = client.get(f"/api/courses/{cid}/documents/{docid}/coverage").json()
    assert coverage["vision"] == 33
    assert document["status"] == "ready"
    assert document["warning"] is None
    assert attempts[before:] == [(27, 2400)]


def _tiny_png():
    buf = io.BytesIO()
    Image.new("RGB", (32, 24), "white").save(buf, "PNG")
    return buf.getvalue()


def test_long_retry_after_halts_requests_but_keeps_cached_content(client, monkeypatch):
    saved = {"n": 0}

    class Response:
        def __init__(self, status, content=""):
            self.status_code = status
            self.headers = {"Retry-After": "90"} if status == 429 else {}
            self.request = httpx.Request("POST", "http://localhost:11434/v1/chat/completions")
            self.content = content
        def raise_for_status(self):
            if self.status_code >= 400:
                raise httpx.HTTPStatusError("upstream", request=self.request, response=self)
        def json(self):
            return {"choices": [{"message": {"content": self.content}}]}

    class Client:
        def __init__(self, *args, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): return False
        def post(self, *args, **kwargs):
            saved["n"] += 1
            return Response(200, "Diagram described") if saved["n"] == 1 else Response(429)

    monkeypatch.setattr(vision.httpx, "Client", Client)
    session = vision.VisionSession("sha1")
    image = _tiny_png()
    assert session.analyze(image, "native", "slide 1") == "Diagram described"
    assert session.analyze(image, "native", "slide 2") is None
    assert "rate-limiting" in session.warning()
    assert session.analyze(image, "native", "slide 3") is None
    assert saved["n"] == 2
    # Cached results remain usable, even after a provider-wide pause.
    assert session.analyze(image, "native", "slide 1") == "Diagram described"
    assert saved["n"] == 2


def test_http_401_is_permanent_without_retrying_or_exposing_error_body(client, monkeypatch):
    seen = []

    class Client:
        def __init__(self, *args, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): return False
        def post(self, url, **kwargs):
            seen.append(url)
            response = httpx.Response(401, request=httpx.Request("POST", url),
                                      text="private-provider-token-should-not-appear")
            response.raise_for_status()

    monkeypatch.setattr(vision.httpx, "Client", Client)
    session = vision.VisionSession("sha2")
    assert session.analyze(_tiny_png(), "text", "slide 1") is None
    assert "HTTP 401" in session.warning()
    assert "private-provider-token" not in session.warning()
    assert session.analyze(_tiny_png(), "text", "slide 2") is None
    assert len(seen) == 1


def test_text_block_response_uses_visible_content_not_hidden_reasoning():
    content, incomplete = vision._extract_response({"choices": [{"message": {
        "reasoning_content": "Private analysis that must not be indexed",
        "content": [{"type": "text", "text": "Price and quantity axis"}]}}]}, "openrouter")
    assert content == "Price and quantity axis"
    assert incomplete is False
