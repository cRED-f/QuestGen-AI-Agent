"""Regression tests for PPTX OCR diagnostics and explicit re-indexing."""
import io
import time

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend import app as service, documents


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("QUESTGEN_DB_PATH", str(tmp_path / "test.sqlite3"))
    monkeypatch.setenv("QUESTGEN_UPLOAD_DIR", str(tmp_path / "files"))
    monkeypatch.delenv("QUESTGEN_EMBED_MODEL", raising=False)
    monkeypatch.delenv("QUESTGEN_VISION_MODEL", raising=False)
    with TestClient(service.app) as test_client:
        yield test_client


def make_deck():
    image = Image.new("RGB", (1200, 560), "white")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = "Bacteria divide by binary fission."
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), Inches(1), Inches(3), width=Inches(5))
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def uploaded(client):
    cid = client.post("/api/courses", json={"name": "Biology"}).json()["id"]
    response = client.post(f"/api/courses/{cid}/documents", files={"file": ("ch2.pptx", make_deck())})
    assert response.status_code == 201
    return cid, response.json()["id"]


def wait(client, cid, docid):
    for _ in range(200):
        result = client.get(f"/api/courses/{cid}/documents/{docid}").json()
        if result["status"] in {"ready", "failed"}:
            return result
        time.sleep(.02)
    raise AssertionError("Document processing did not finish")


def test_plain_text_remains_indexed_when_vision_unavailable(client, monkeypatch):
    monkeypatch.setenv("QUESTGEN_VISION_ENABLED", "false")
    cid, docid = uploaded(client)
    result = wait(client, cid, docid)
    assert result["status"] == "ready"
    assert "Vision unavailable" in result["warning"]
    assert client.get(f"/api/courses/{cid}/search", params={"query": "binary fission"}).json()


def test_reprocess_uses_selected_model_and_replaces_old_chunks(client, monkeypatch):
    from backend import vision
    monkeypatch.setenv("QUESTGEN_VISION_ENABLED", "true")
    client.put("/api/settings", json={"provider": "ollama", "model": "fake-vision-model"})
    # Deterministic slide-rendering fixture: exercise the ingestion contract
    # without introducing LibreOffice process startup flakiness into unit tests.
    from backend import ingestion
    import fitz
    def renderer(source, output):
        target = output / (source.stem + ".pdf")
        document = fitz.open()
        document.new_page(width=900, height=550)
        document.save(str(target))
        document.close()
        return target
    monkeypatch.setattr(ingestion, "_render_office", renderer)
    statements = []
    def first(self, image, native, label):
        statements.append(label)
        return "Initial visual labels: one two three"
    monkeypatch.setattr(vision.VisionSession, "analyze", first)
    cid, docid = uploaded(client)
    assert wait(client, cid, docid)["status"] == "ready"
    first_count = client.get(f"/api/courses/{cid}/documents/{docid}/coverage").json()
    assert first_count["vision"] == 1
    def updated(self, image, native, label):
        return "Updated visual evidence: genome replication"
    monkeypatch.setattr(vision.VisionSession, "analyze", updated)
    response = client.post(f"/api/courses/{cid}/documents/{docid}/reprocess")
    assert response.status_code == 202
    assert wait(client, cid, docid)["status"] == "ready"
    found = client.get(f"/api/courses/{cid}/search", params={"query": "genome replication"}).json()
    assert found and found[0]["page"] == 1
    assert not client.get(f"/api/courses/{cid}/search", params={"query": "Initial"}).json()
    assert client.post(f"/api/courses/other/documents/{docid}/reprocess").status_code == 404


def test_ingestion_status_never_returns_model_secrets(client):
    response = client.get("/api/ingestion/status")
    assert response.status_code == 200
    data = response.json()
    assert "renderer_available" in data and "message" in data
    assert "api_key" not in response.text
    assert client.get("/api/ocr/status").status_code == 404
