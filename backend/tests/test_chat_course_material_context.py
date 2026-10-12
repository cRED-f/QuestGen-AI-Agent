"""Regression: uploaded PowerPoints must be visible to broad assessment requests."""
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from backend import app as service


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("QUESTGEN_DB_PATH", str(tmp_path / "db.sqlite3"))
    monkeypatch.setenv("QUESTGEN_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.delenv("QUESTGEN_EMBED_MODEL", raising=False)
    with TestClient(service.app) as api:
        api.put("/api/settings", json={"provider": "ollama", "model": "test-model"})
        yield api


def upload_deck(client, course_id: str, name: str, topic: str) -> str:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    textbox.text = topic
    data = io.BytesIO()
    presentation.save(data)
    response = client.post(f"/api/courses/{course_id}/documents", files={"file": (name, data.getvalue())})
    assert response.status_code == 201
    doc_id = response.json()["id"]
    for _ in range(100):
        document = client.get(f"/api/courses/{course_id}/documents/{doc_id}").json()
        if document["status"] in {"ready", "failed"}:
            break
        time.sleep(.03)
    assert document["status"] == "ready", document
    assert document["chunk_count"] > 0
    return doc_id


def events(response):
    assert response.status_code == 200, response.text
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]


def test_uploaded_pptx_available_after_hi_and_on_generic_question_paper_request(client, monkeypatch):
    course_id = client.get("/api/courses").json()[0]["id"]
    first = upload_deck(client, course_id, "Ch 3 PPT.pptx", "Stacks use a last-in-first-out structure")
    second = upload_deck(client, course_id, "Chapter 02 PPT.pptx", "Trees contain nodes with parent-child links")
    thread_id = client.post("/api/threads", json={"course_id": course_id}).json()["id"]

    prompts = []
    async def fake_generate(_settings, history):
        prompts.append(history)
        yield "Ready to use the uploaded material."
    monkeypatch.setattr(service, "generate_tokens", fake_generate)

    hello = events(client.post(f"/api/threads/{thread_id}/stream", json={"content": "hi"}))
    assert next(event for event in hello if event["type"] == "sources")["sources"] == []
    # A greeting need not scan slide text, but it must not make the model think
    # that the library contains no files.
    hello_context = "\n".join(item["content"] for item in prompts[-1])
    assert "Ch 3 PPT.pptx" in hello_context
    assert "Chapter 02 PPT.pptx" in hello_context
    assert "ready" in hello_context

    paper = events(client.post(f"/api/threads/{thread_id}/stream", json={"content": "I want to make question paper"}))
    sources = next(event for event in paper if event["type"] == "sources")["sources"]
    assert {source["document_id"] for source in sources} == {first, second}
    assert {source["filename"] for source in sources} == {"Ch 3 PPT.pptx", "Chapter 02 PPT.pptx"}
    assert all(source["page"] == 1 for source in sources)
    assert any("Stacks use" in item["content"] for item in prompts[-1])
    assert any("Trees contain" in item["content"] for item in prompts[-1])


def test_selected_materials_remain_strictly_scoped(client, monkeypatch):
    cid = client.get("/api/courses").json()[0]["id"]
    first = upload_deck(client, cid, "physics.pptx", "Newton formulated laws of motion")
    second = upload_deck(client, cid, "economics.pptx", "Market equilibrium relates price and demand")
    tid = client.post("/api/threads", json={"course_id": cid}).json()["id"]
    observed = []
    async def fake_generate(_settings, history):
        observed.append(history)
        yield "Ready."
    monkeypatch.setattr(service, "generate_tokens", fake_generate)

    reply = events(client.post(f"/api/threads/{tid}/stream", json={
        "content": "Please create an exam question paper", "document_ids": [first]
    }))
    sources = next(event for event in reply if event["type"] == "sources")["sources"]
    assert sources and {s["document_id"] for s in sources} == {first}
    model_context = "\n".join(item["content"] for item in observed[-1])
    assert "physics.pptx" in model_context
    assert "economics.pptx" not in model_context
    assert second not in {s["document_id"] for s in sources}


def test_generic_question_paper_not_grounded_on_another_course(client, monkeypatch):
    course_a = client.get("/api/courses").json()[0]["id"]
    other_course = client.post("/api/courses", json={"name": "Other"}).json()["id"]
    upload_deck(client, other_course, "secret-course.pptx", "This is a separate university syllabus")
    tid = client.post("/api/threads", json={"course_id": course_a}).json()["id"]
    observed = []
    async def fake_generate(_settings, history):
        observed.append(history)
        yield "No ready materials in the current course."
    monkeypatch.setattr(service, "generate_tokens", fake_generate)
    response = events(client.post(f"/api/threads/{tid}/stream", json={"content": "Make question paper"}))
    assert next(event for event in response if event["type"] == "sources")["sources"] == []
    assert "secret-course.pptx" not in str(observed[-1])


def upload_multislide_deck(client, course_id: str, name: str, count: int = 47) -> str:
    presentation = Presentation()
    for slide_number in range(1, count + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(2))
        textbox.text = f"Slide {slide_number}: topic-specific lecture fact {slide_number} — explanation from the instructor."
    data = io.BytesIO()
    presentation.save(data)
    response = client.post(f"/api/courses/{course_id}/documents", files={"file": (name, data.getvalue())})
    assert response.status_code == 201
    doc_id = response.json()["id"]
    for _ in range(200):
        doc = client.get(f"/api/courses/{course_id}/documents/{doc_id}").json()
        if doc["status"] in {"ready", "failed"}:
            break
        time.sleep(0.03)
    assert doc["status"] == "ready", doc
    assert doc["pages"] == count
    return doc_id


def test_explain_chapter_3_reads_full_deck_not_only_lexical_hits(client, monkeypatch):
    course_id = client.get("/api/courses").json()[0]["id"]
    first = upload_multislide_deck(client, course_id, "Ch 3 PPT.pptx")
    second = upload_deck(client, course_id, "Chapter 02 PPT.pptx", "Unrelated chapter 2 material")
    thread_id = client.post("/api/threads", json={"course_id": course_id}).json()["id"]
    inputs = []

    async def fake_generate(_settings, history, **kwargs):
        inputs.append((history, kwargs))
        yield "The material presents 47 individual slides [47]."

    monkeypatch.setattr(service, "generate_tokens", fake_generate)
    streamed = events(client.post(f"/api/threads/{thread_id}/stream", json={
        "content": "explain the chapter 3 ppt", "document_ids": [first, second],
    }))
    sources = next(e for e in streamed if e["type"] == "sources")["sources"]
    assert len(sources) == 47  # 47 slides, not the six best lexical matches
    assert {source["page"] for source in sources} == set(range(1, 48))
    assert {source["document_id"] for source in sources} == {first}
    prompt = "\n".join(item["content"] for item in inputs[-1][0])
    assert inputs[-1][1]["max_tokens"] == 3200
    assert "Slide 3: topic-specific" in prompt
    assert "Slide 20: topic-specific" in prompt
    assert "Slide 47: topic-specific" in prompt
    assert "47" in prompt and "Whole-document explanation" in prompt
    assert "Unrelated chapter 2 material" not in prompt
    stored = client.get(f"/api/threads/{thread_id}/messages").json()[-1]
    assert len(stored["sources"]) == 47
    assert stored["sources"][-1]["page"] == 47


def test_generic_question_stays_small_and_ambiguous_presentation_does_not_mix_files(client, monkeypatch):
    cid = client.get("/api/courses").json()[0]["id"]
    upload_multislide_deck(client, cid, "Ch 3 PPT.pptx", 12)
    upload_deck(client, cid, "Chapter 02 PPT.pptx", "An entirely different lesson")
    tid = client.post("/api/threads", json={"course_id": cid}).json()["id"]
    calls = []
    async def fake_generate(_settings, history):
        calls.append(history)
        yield "Response."
    monkeypatch.setattr(service, "generate_tokens", fake_generate)

    normal = events(client.post(f"/api/threads/{tid}/stream", json={"content": "What is topic-specific?"}))
    assert len(next(e for e in normal if e["type"] == "sources")["sources"]) <= 6
    broad = events(client.post(f"/api/threads/{tid}/stream", json={"content": "Explain the PPT"}))
    assert next(e for e in broad if e["type"] == "sources")["sources"] == []
    assert "Whole-document explanation" not in str(calls[-1])


def test_named_deck_request_never_fetches_unselected_deck(client, monkeypatch):
    cid = client.get("/api/courses").json()[0]["id"]
    other = client.post("/api/courses", json={"name": "Other"}).json()["id"]
    upload_deck(client, cid, "Ch 3 PPT.pptx", "Lecture on microeconomic markets")
    selected = upload_deck(client, cid, "Chapter 02 PPT.pptx", "Chapter two property laws")
    upload_deck(client, other, "Ch 3 PPT.pptx", "Private course source")
    tid = client.post("/api/threads", json={"course_id": cid}).json()["id"]
    calls = []
    async def fake_generate(_settings, history, **_kwargs):
        calls.append(str(history))
        yield "Response."
    monkeypatch.setattr(service, "generate_tokens", fake_generate)
    result = events(client.post(f"/api/threads/{tid}/stream", json={
        "content": "Explain the chapter 3 ppt", "document_ids": [selected],
    }))
    # Explicit attachment restriction takes precedence even over a different named deck.
    sources = next(e for e in result if e["type"] == "sources")["sources"]
    assert all(source["document_id"] == selected for source in sources)
    assert "Private course source" not in calls[-1]
    assert "microeconomic markets" not in calls[-1]
