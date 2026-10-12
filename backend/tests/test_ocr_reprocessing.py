"""Regression tests for PPTX OCR diagnostics and explicit re-indexing."""
import io
import time

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend import app as service, documents


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("QUESTGEN_DB_PATH", str(tmp_path / "test.sqlite3"))
    monkeypatch.setenv("QUESTGEN_UPLOAD_DIR", str(tmp_path / "files"))
    monkeypatch.delenv("QUESTGEN_EMBED_MODEL", raising=False)
    monkeypatch.delenv("QUESTGEN_VISION_MODEL", raising=False)
    with TestClient(service.app) as test_client:
        yield test_client


def make_deck():
    image = Image.new("RGB", (1200, 560), "white")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = "Bacteria divide by binary fission."
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), Inches(1), Inches(3), width=Inches(5))
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def uploaded(client):
    cid = client.post("/api/courses", json={"name": "Biology"}).json()["id"]
    response = client.post(f"/api/courses/{cid}/documents", files={"file": ("ch2.pptx", make_deck())})
    assert response.status_code == 201
    return cid, response.json()["id"]


def wait(client, cid, docid):
    for _ in range(200):
        result = client.get(f"/api/courses/{cid}/documents/{docid}").json()
        if result["status"] in {"ready", "failed"}:
            return result
        time.sleep(.02)
    raise AssertionError("Document processing did not finish")


def test_ocr_no_text_is_not_reported_as_missing_tesseract(client, monkeypatch):
    monkeypatch.setattr(documents, "image_content", lambda *_args, **_kw: "")
    cid, docid = uploaded(client)
    result = wait(client, cid, docid)
    assert result["status"] == "ready"
    assert "no OCR-readable text" in result["warning"]
    assert "Install OCR" not in result["warning"]
    assert "Tesseract OCR is not accessible" not in result["warning"]
    assert client.get(f"/api/courses/{cid}/search", params={"query": "binary fission"}).json()


def test_reprocess_after_ocr_becomes_available(client, monkeypatch):
    def missing(*args, **kwargs):
        raise documents.OCRUnavailableError("Tesseract OCR is not accessible")
    monkeypatch.setattr(documents, "image_content", missing)
    cid, docid = uploaded(client)
    result = wait(client, cid, docid)
    assert "OCR is not accessible" in result["warning"]
    monkeypatch.setattr(documents, "image_content", lambda *_a, **_kw: "Visible text (OCR): genome replication")
    response = client.post(f"/api/courses/{cid}/documents/{docid}/reprocess")
    assert response.status_code == 202
    restored = wait(client, cid, docid)
    assert restored["warning"] is None
    assert restored["status"] == "ready"
    matches = client.get(f"/api/courses/{cid}/search", params={"query": "genome replication"}).json()
    assert matches and matches[0]["page"] == 1
    assert client.post(f"/api/courses/unrelated/documents/{docid}/reprocess").status_code == 404


def test_ocr_status_does_not_disclose_executable_path(client, monkeypatch):
    monkeypatch.setenv("QUESTGEN_TESSERACT_CMD", "C:/Private/Tools/tesseract.exe")
    monkeypatch.setattr("pytesseract.get_tesseract_version", lambda: "5.5.0")
    response = client.get("/api/ocr/status")
    assert response.status_code == 200
    result = response.json()
    assert result["available"] is True
    assert result["source"] == "configured"
    assert "C:/Private" not in response.text
    assert result["version"] == "5.5.0"


def test_ocr_status_missing_executable(client, monkeypatch):
    import pytesseract
    monkeypatch.setattr("pytesseract.get_tesseract_version", lambda: (_ for _ in ()).throw(pytesseract.TesseractNotFoundError()))
    status = client.get("/api/ocr/status").json()
    assert status["available"] is False
    assert "QUESTGEN_TESSERACT_CMD" in status["message"]


def test_tesseract_override_takes_priority(monkeypatch):
    monkeypatch.setenv("QUESTGEN_TESSERACT_CMD", '"C:/Program Files/Tesseract-OCR/tesseract.exe"')
    path, source = documents.tesseract_command()
    assert path == "C:/Program Files/Tesseract-OCR/tesseract.exe"
    assert source == "configured"
