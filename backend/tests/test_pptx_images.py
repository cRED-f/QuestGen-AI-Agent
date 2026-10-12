"""PPTX/image ingestion, OCR/vision fallback, and source attribution regressions."""
import io
import os
import time
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from backend import app as service, documents


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("QUESTGEN_DB_PATH", str(tmp_path / "db.sqlite3"))
    monkeypatch.setenv("QUESTGEN_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.delenv("QUESTGEN_EMBED_MODEL", raising=False)
    monkeypatch.delenv("QUESTGEN_VISION_MODEL", raising=False)
    with TestClient(service.app) as client:
        yield client


def course(client):
    return client.post("/api/courses", json={"name": "Course"}).json()["id"]


def wait(client, cid, docid):
    for _ in range(200):
        result = client.get(f"/api/courses/{cid}/documents/{docid}").json()
        if result["status"] in ("ready", "failed"):
            return result
        time.sleep(0.02)
    raise AssertionError("Document ingestion did not finish")


def upload(client, cid, data, filename):
    return client.post(f"/api/courses/{cid}/documents", files={"file": (filename, data)})


def png_bytes():
    image = Image.new("RGB", (640, 160), "white")
    ImageDraw.Draw(image).text((20, 50), "Recursion uses a base case", fill="black")
    data = io.BytesIO()
    image.save(data, format="PNG")
    return data.getvalue()


def pptx_bytes(with_image=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = "Stacks obey last-in-first-out ordering."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "LIFO"
    table.cell(0, 1).text = "stack"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = "A binary heap supports priority queues."
    if with_image:
        slide2.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(3), width=Inches(3))
    data = io.BytesIO()
    presentation.save(data)
    return data.getvalue()


def test_pptx_slide_search_citation_and_download(client):
    cid, second = course(client), course(client)
    response = upload(client, cid, pptx_bytes(), "module.pptx")
    assert response.status_code == 201, response.text
    doc = wait(client, cid, response.json()["id"])
    assert doc["status"] == "ready", doc
    assert doc["pages"] == 2
    assert doc["chunk_count"] >= 2
    found = client.get(f"/api/courses/{cid}/search", params={"query": "priority queues binary heap"}).json()
    assert len(found) >= 1
    assert found[0]["page"] == 2 and found[0]["extension"] == ".pptx"
    assert found[0]["source_url"].endswith("/file")
    assert client.get(f"/api/courses/{second}/search", params={"query": "heap"}).json() == []
    file_reply = client.get(f"/api/courses/{cid}/documents/{doc['id']}/file")
    assert file_reply.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert file_reply.content.startswith(b"PK")
    assert client.delete(f"/api/courses/{cid}/documents/{doc['id']}").status_code == 204
    assert client.get(f"/api/courses/{cid}/search", params={"query": "heap"}).json() == []


def test_slide_embedded_image_ocr(client, monkeypatch):
    monkeypatch.setattr(documents, "image_content", lambda blob, **kwargs: "OCR diagram: B-tree indexing")
    cid = course(client)
    reply = upload(client, cid, pptx_bytes(with_image=True), "slides.pptx")
    assert wait(client, cid, reply.json()["id"])["status"] == "ready"
    found = client.get(f"/api/courses/{cid}/search", params={"query": "indexing"}).json()
    assert found and found[0]["page"] == 2
    assert "Image on this slide" in found[0]["snippet"]


def test_png_and_jpeg_image_retrieval_with_ocr(client, monkeypatch):
    monkeypatch.setattr("pytesseract.image_to_string", lambda *_args, **_kwargs: "Mitosis has four named phases")
    cid = course(client)
    data = png_bytes()
    reply = upload(client, cid, data, "biology.png")
    assert wait(client, cid, reply.json()["id"])["status"] == "ready"
    found = client.get(f"/api/courses/{cid}/search", params={"query": "Mitosis"}).json()
    assert found and found[0]["extension"] == ".png"
    assert "Mitosis" in found[0]["snippet"]
    image_reply = client.get(f"/api/courses/{cid}/documents/{reply.json()['id']}/file")
    assert image_reply.headers["content-type"].startswith("image/png")
    assert image_reply.headers["content-disposition"].startswith("inline")
    image = Image.open(io.BytesIO(data))
    buf = io.BytesIO()
    image.save(buf, "JPEG")
    assert upload(client, cid, buf.getvalue(), "lecture.jpeg").status_code == 201
    assert upload(client, cid, data, "incorrect.jpg").status_code == 422


def test_invalid_corrupt_and_bomb_pptx(client):
    cid = course(client)
    assert upload(client, cid, b"not-pptx", "fake.pptx").status_code == 422
    assert upload(client, cid, b"not-png", "fake.png").status_code == 422
    data = io.BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as zip_file:
        zip_file.writestr("ppt/presentation.xml", b"a" * (101 * 1024 * 1024))
        zip_file.writestr("[Content_Types].xml", b"stub")
    assert upload(client, cid, data.getvalue(), "bomb.pptx").status_code == 422


def test_no_readable_text_reports_failure(client, monkeypatch):
    monkeypatch.setattr("pytesseract.image_to_string", lambda *_args, **_kwargs: "   ")
    cid = course(client)
    reply = upload(client, cid, png_bytes(), "diagram.png")
    doc = wait(client, cid, reply.json()["id"])
    assert doc["status"] == "failed"
    assert "VISION_MODEL" in doc["error"]
    assert client.get(f"/api/courses/{cid}/search", params={"query": "diagram"}).json() == []


def test_missing_tesseract_gives_actionable_error(monkeypatch):
    from pytesseract import pytesseract
    def raise_missing(*args, **kwargs):
        raise pytesseract.TesseractNotFoundError()
    monkeypatch.setattr("pytesseract.image_to_string", raise_missing)
    with pytest.raises(ValueError, match="Tesseract OCR"):
        documents.image_content(png_bytes(), extension=".png")


def test_vision_descriptions_for_diagrams_without_ocr(monkeypatch, client):
    from pytesseract import pytesseract
    monkeypatch.setattr("pytesseract.image_to_string", lambda *_a, **_kw: "")
    monkeypatch.setenv("QUESTGEN_VISION_MODEL", "llava-local")
    called = []
    class FakeReply:
        def raise_for_status(self):
            return None
        def json(self):
            return {"message": {"content": "The flow diagram includes two branches: pass and fail."}}
    class FakeHttpClient:
        def __init__(self, *args, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def post(self, url, **kwargs):
            called.append((url, kwargs["json"]))
            return FakeReply()
    cid = course(client)
    monkeypatch.setattr(documents.httpx, "Client", FakeHttpClient)
    response = upload(client, cid, png_bytes(), "flow.png")
    assert wait(client, cid, response.json()["id"])["status"] == "ready"
    found = client.get(f"/api/courses/{cid}/search", params={"query": "branches fail"}).json()
    assert found
    assert len(called) == 1
    assert called[0][0].endswith("/api/chat")
    assert called[0][1]["messages"][0]["images"]


def test_pptx_keeps_text_when_embedded_image_unreadable(client, monkeypatch):
    def ocr_unavailable(*args, **kwargs):
        raise documents.OCRUnavailableError("Tesseract OCR is not accessible to the FastAPI backend")
    monkeypatch.setattr(documents, "image_content", ocr_unavailable)
    cid = course(client)
    response = upload(client, cid, pptx_bytes(with_image=True), "mixed.pptx")
    indexed = wait(client, cid, response.json()["id"])
    assert indexed["status"] == "ready"
    assert "Tesseract OCR is not accessible" in indexed["warning"]
    assert client.get(f"/api/courses/{cid}/search", params={"query": "priority queues"}).json()


def test_repeated_slide_pictures_are_processed_once(monkeypatch):
    data = png_bytes()
    p = Presentation()
    for _ in range(3):
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(data), Inches(1), Inches(1), width=Inches(3))
    stream = io.BytesIO()
    p.save(stream)
    calls = []
    def ocr(blob, **kwargs):
        calls.append(blob)
        return "Cell division has mitosis stages"
    monkeypatch.setattr(documents, "image_content", ocr)
    import tempfile
    with tempfile.TemporaryDirectory() as folder:
        path = __import__("pathlib").Path(folder) / "images.pptx"
        path.write_bytes(stream.getvalue())
        segments = documents.extract_segments(path, ".pptx")
    assert len(segments) == 3
    assert len(calls) == 1
